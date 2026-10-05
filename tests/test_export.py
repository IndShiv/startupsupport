"""CSV and Excel exports of the filtered lists."""

import csv
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from crm.exports import safe_text
from crm.models import Coach, ExportLog, Startup, Tag

pytestmark = pytest.mark.django_db


@pytest.fixture
def roeland_client(client, make_user):
    client.force_login(make_user("roeland", "Coach", coach=Coach.objects.get(first_name="Roeland")))
    return client


@pytest.fixture
def admin_client(client, make_user):
    client.force_login(make_user("boss", "Admin"))
    return client


def read_csv(response):
    text = response.content.decode("utf-8")
    assert text.startswith("﻿")  # BOM for Excel
    return list(csv.DictReader(io.StringIO(text[1:])))


def read_xlsx(response):
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)], wb


def test_requires_staff(client, make_user, reference):
    assert client.get("/staff/export/startups.csv").status_code == 302
    client.force_login(make_user("nobody"))
    assert client.get("/staff/export/startups.csv").status_code == 403


def test_unknown_list_or_format(admin_client):
    assert admin_client.get("/staff/export/coaches.csv").status_code == 404
    assert admin_client.get("/staff/export/startups.pdf").status_code == 404


def test_students_csv(admin_client, register):
    register(first_name="Zoë", last_name="Ünal", phone="0612345678")
    response = admin_client.get("/staff/export/students.csv")
    assert response["Content-Type"].startswith("text/csv")
    assert 'filename="buss-students-' in response["Content-Disposition"]
    assert response["Cache-Control"] == "no-store"
    [row] = read_csv(response)
    assert (row["First name"], row["Last name"], row["Student number"], row["Ideas"]) == ("Zoë", "Ünal", "234567", "1")
    assert row["Domain"] == "Games" and row["Graduation track"] == "No"


def test_startups_xlsx(admin_client, register, graduation_data):
    reg = register(**graduation_data(grad_approval="not_yet"))
    Startup.objects.filter(pk=reg.startup_id).update(name="Polder Games", has_paying_customers=True)
    response = admin_client.get("/staff/export/startups.xlsx")
    rows, wb = read_xlsx(response)
    [row] = rows
    assert row["Startup name"] == "Polder Games"
    assert row["Founders"] == "Sanne de Vries"
    assert row["Paying customers"] == "Yes"
    assert row["Programme approval"] == "No, not yet"
    assert isinstance(row["Hand-in date"], (date,))
    ws = wb.worksheets[0]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    about = dict(wb["About this export"].iter_rows(values_only=True, max_col=2))
    assert about["Rows"] == 1 and "personal data" in about["Privacy"]


def test_export_follows_filters(admin_client, register):
    a = register(student_number="100001", email="a@buas.nl")
    b = register(student_number="100002", email="b@buas.nl")
    tag = Tag.objects.create(name="impact")
    a.startup.tags.add(tag)
    rows = read_csv(admin_client.get(f"/staff/export/startups.csv?tags={tag.pk}"))
    assert [int(r["ID"]) for r in rows] == [a.startup_id]
    assert len(read_csv(admin_client.get("/staff/export/startups.csv"))) == 2
    rows, wb = read_xlsx(admin_client.get(f"/staff/export/startups.xlsx?tags={tag.pk}"))
    about = dict(wb["About this export"].iter_rows(values_only=True, max_col=2))
    assert about["Filter: Tags"] == "impact"


def test_coach_export_defaults_to_own_caseload(roeland_client, register):
    mine = register(student_number="100001", email="a@buas.nl")
    register(student_number="100002", email="b@buas.nl")
    Startup.objects.filter(pk=mine.startup_id).update(assigned_coach=Coach.objects.get(first_name="Roeland"))
    assert [int(r["ID"]) for r in read_csv(roeland_client.get("/staff/export/startups.csv"))] == [mine.startup_id]
    assert len(read_csv(roeland_client.get("/staff/export/startups.csv?scope=all"))) == 2


def test_every_export_is_logged(admin_client, register):
    register()
    admin_client.get("/staff/export/students.xlsx?q=sanne")
    log = ExportLog.objects.get()
    assert (log.kind, log.file_format, log.row_count, log.user.username) == ("students", "xlsx", 1, "boss")
    assert log.filters["Search"] == "sanne"


@pytest.mark.parametrize("value, csv_expected, xlsx_expected", [
    ("=HYPERLINK(\"http://evil\")", "'=HYPERLINK(\"http://evil\")", "'=HYPERLINK(\"http://evil\")"),
    ("+31 6 12345678", "'+31 6 12345678", "+31 6 12345678"),
    ("-10", "'-10", "-10"),
    ("@SUM(A1)", "'@SUM(A1)", "@SUM(A1)"),
    ("Normal text", "Normal text", "Normal text"),
    (42, 42, 42),
])
def test_formula_injection_is_neutralised(value, csv_expected, xlsx_expected):
    assert safe_text(value) == csv_expected
    assert safe_text(value, xlsx=True) == xlsx_expected


def test_malicious_form_input_is_safe_in_both_formats(admin_client, register):
    register(first_name="=cmd|' /C calc'!A0", description="=HYPERLINK(\"http://evil.example\",\"click\")")
    [row] = read_csv(admin_client.get("/staff/export/students.csv"))
    assert row["First name"].startswith("'=")
    rows, wb = read_xlsx(admin_client.get("/staff/export/startups.xlsx"))
    cell = next(c for c in wb.worksheets[0][2] if isinstance(c.value, str) and "HYPERLINK" in c.value)
    assert cell.data_type == "s" and cell.value.startswith("'=")


def test_export_buttons_keep_filters(admin_client, register):
    register()
    content = admin_client.get("/staff/startups/?scope=all&paying=yes").content.decode()
    assert "/staff/export/startups.xlsx?scope=all&amp;paying=yes" in content
