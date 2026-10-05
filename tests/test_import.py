"""Spreadsheet import, using a fake workbook shaped like the BUSS coach overview."""

import io
from datetime import datetime

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from crm import importer
from crm.models import Coach, Founder, FollowUp, GraduationTrack, ImportBatch, Registration, Startup, Student

pytestmark = pytest.mark.django_db

HEADER = ["Student (GROEN = actief, ORANJE = inactief, ROOD = afgerond,  BLAUW = medewerker)", "student number",
          "studyprogram", "BUSS coach", "Graduation own company (GOC)", "Company name", "Description company",
          "Deadline thesis", "Graduation supervisor Academy", "Date end evaluation coach track ", "status (coaching) track"]
GREEN, BLUE, RED = "FF92D050", "FF00B0F0", "FFFF0000"
ROWS = [
    # name, number, program, coach, goc, company, description, deadline, supervisor, end, status, colour
    ("Anna Jansen", 111111, "Logistics", "Erik", "no", "Cargo Bikes BV", "Last-mile delivery by cargo bike", None, None, None, "Coaching every month", GREEN),
    ("Bram de Vries", "222222", "Data Science &AI", "Tijs", "YES", None, "Litter detection with drones", datetime(2027, 1, 15), "Dr. Smit", None, None, GREEN),
    ("Chris Peters", None, "Buas Emplyee", "Hans", None, None, "Side business in coaching", None, None, None, None, BLUE),
    ("Dana Bos & Eva Kok", "333333 444444", "Tourism", "Niki", "No", None, "Bike tours", None, None, None, None, None),
    ("Fleur Visser", "Alumnus - 555555", "ALUMNUS MMI", "Marc", "yes", None, "Media studio", "?", "?", None, None, RED),
    ("Gijs Mulder", "Employee", "Masters Media", "Ben", "NO", None, "Podcast network", None, None, None, None, GREEN),
]


def workbook_bytes(rows=ROWS, sheet="2026-2027", extra_sheet=True):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(HEADER)
    for row in rows:
        ws.append(list(row[:-1]))
        if row[-1]:
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=row[-1])
    ws.append([None] * len(HEADER))  # empty rows are ignored
    if extra_sheet:
        wb.create_sheet("2025-2026").append(HEADER)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def batch(reference):
    header, rows = importer.parse_sheet(io.BytesIO(workbook_bytes()), "2026-2027")
    b = ImportBatch.objects.create(filename="Coach_overview.xlsx", sheet="2026-2027", header=header, rows=rows)
    b.mapping = importer.default_mapping(header, rows, b.sheet)
    b.save()
    return b


# -- parsing and guessing ---------------------------------------------------------------------

def test_parse_sheet_reads_values_and_colours(reference):
    header, rows = importer.parse_sheet(io.BytesIO(workbook_bytes()), "2026-2027")
    assert header[1] == "student number" and len(rows) == len(ROWS)
    assert rows[0]["cells"][1] == "111111"  # numbers become text without ".0"
    assert rows[1]["cells"][7].startswith("2027-01-15")
    assert [r["fill"] for r in rows] == [GREEN, GREEN, BLUE, "", RED, GREEN]
    assert importer.sheet_names(io.BytesIO(workbook_bytes())) == ["2026-2027", "2025-2026"]


def test_guess_columns(reference):
    columns = importer.guess_columns(HEADER)
    named = {t: HEADER[i] for t, i in columns.items() if i is not None}
    assert named["full_name"].startswith("Student (GROEN")
    assert named["student_number"] == "student number"
    assert named["program"] == "studyprogram"
    assert named["coach"] == "BUSS coach"
    assert named["goc"].startswith("Graduation own company")
    assert named["company_name"] == "Company name" and named["description"] == "Description company"
    assert named["hand_in"] == "Deadline thesis" and named["notes"] == "status (coaching) track"
    assert columns["email"] is None and columns["submitted_at"] is None


def test_guess_columns_for_a_forms_export(reference):
    header = ["ID", "Start time", "Completion time", "Email", "Name", "First name", "Last name", "Student number", "Phone number", "Domain", "Study year"]
    columns = importer.guess_columns(header)
    assert header[columns["response_id"]] == "ID"
    assert header[columns["submitted_at"]] == "Completion time"
    assert header[columns["first_name"]] == "First name" and header[columns["last_name"]] == "Last name"
    assert columns["full_name"] is None


@pytest.mark.parametrize("value, domain, year", [
    ("Logistics", "Logistics", ""),
    ("Data Science &AI", "Data Science & AI", ""),
    ("Buas Emplyee", "I am an employee", "Not applicable"),
    ("BUas employees", "I am an employee", "Not applicable"),
    ("Masters Media", "Media", "BUas Master"),
    ("Master SEM", "Leisure & Events", "BUas Master"),
    ("ALUMNUS MMI", "Media", "Not applicable"),
    ("Alumni", importer.UNKNOWN_DOMAIN, "Not applicable"),
    ("", importer.UNKNOWN_DOMAIN, ""),
])
def test_guess_program(reference, value, domain, year):
    assert importer.guess_program(value) == (domain, year)


def test_guess_coach_and_yes_no(reference):
    assert importer.guess_coach("Erik") == str(Coach.objects.get(first_name="Erik").pk)
    assert importer.guess_coach("Niki") == ""
    assert [importer.guess_yes_no(v) for v in ["YES", "yes", "Ja", "NO", "no", "nee", "", "?"]] == ["yes", "yes", "yes", "no", "no", "no", "unknown", "unknown"]


def test_default_date_from_sheet_name():
    assert importer.default_date_for("2026-2027").isoformat() == "2026-09-01"
    assert importer.default_date_for("2024 - 2025").isoformat() == "2024-09-01"


# -- preview ------------------------------------------------------------------------------------

def plan_by_row(batch):
    return {p.row: p for p in importer.build_plan(batch)}


def test_preview(batch):
    plan = plan_by_row(batch)
    anna = plan[2]
    assert anna.action == "new" and anna.people[0]["first_name"] == "Anna" and anna.people[0]["student_number"] == "111111"
    assert anna.domain == "Logistics" and anna.year == importer.UNKNOWN_YEAR and anna.coach.first_name == "Erik"
    assert anna.stage == "Coaching" and anna.date_estimated and anna.submitted_at.date().isoformat() == "2026-09-01"
    bram = plan[3]
    assert bram.goc == "yes" and bram.year == "Year 4 and graduating within own company" and str(bram.hand_in) == "2027-01-15"
    chris = plan[4]
    assert chris.domain == "I am an employee" and chris.people[0]["student_number"] == ""
    team = plan[5]
    assert [(p["first_name"], p["last_name"], p["student_number"]) for p in team.people] == [("Dana", "Bos", "333333"), ("Eva", "Kok", "444444")]
    assert any("Niki" in w for w in team.warnings) and "status unknown (imported)" in team.tags
    fleur = plan[6]
    assert fleur.people[0]["student_number"] == "555555" and fleur.stage == "Alumni"
    assert "graduation details missing" in fleur.tags  # "?" is not a date
    gijs = plan[7]
    assert gijs.year == "BUas Master" and any("No 6-digit student number" in w for w in gijs.warnings)


def test_preview_links_known_students(batch, register):
    known = register(student_number="111111", email="anna@buas.nl")
    p = plan_by_row(batch)[2]
    assert p.action == "link" and p.people[0]["existing"] == known.student


def test_preview_does_not_save_anything(batch):
    importer.build_plan(batch)
    assert not Student.objects.exists() and not Registration.objects.exists()


# -- running the import ---------------------------------------------------------------------------

def test_run_import(batch, make_user):
    admin = make_user("boss", "Admin")
    counts = importer.run_import(batch, admin)
    assert counts == {"new": 6, "link": 0, "skip": 0, "error": 0, "students_created": 7}
    assert Registration.objects.filter(source="import").count() == 6
    reg = Registration.objects.get(student__student_number="111111")
    assert reg.submitted_at_estimated and reg.intake_deadline is None and reg.created_by == admin
    assert reg.answers["import"]["values"]["BUSS coach"] == "Erik"
    assert reg.startup.name == "Cargo Bikes BV" and reg.startup.stage.name == "Coaching"
    assert reg.startup.activities.get().body.endswith("Coaching every month")
    team = Startup.objects.get(description="Bike tours")
    assert team.founders.count() == 2 and team.assigned_coach is None
    track = GraduationTrack.objects.get()
    assert track.student.student_number == "222222" and track.approval == "yes" and track.supervisor_name == "Dr. Smit"
    employee = Student.objects.get(first_name="Chris")
    assert employee.email == "" and employee.domain.is_employee
    # Row data is not kept after the import.
    batch.refresh_from_db()
    assert batch.status == "imported" and batch.rows == [] and batch.result["new"] == 6


def test_imports_stay_out_of_queue_and_monthly_figures(batch, make_user):
    from crm import dashboard
    importer.run_import(batch, make_user("boss", "Admin"))
    assert not Registration.objects.awaiting_intake().exists()
    period = dashboard.get_period("all")
    assert dashboard.kpis(period)["registrations"] == 0
    assert dashboard.caseload()[0]["count"] >= 1  # but they do count as caseload


def test_running_twice_skips_imported_rows(batch, make_user):
    importer.run_import(batch, make_user("boss", "Admin"))
    header, rows = importer.parse_sheet(io.BytesIO(workbook_bytes()), "2026-2027")
    again = ImportBatch.objects.create(filename="x.xlsx", sheet="2026-2027", header=header, rows=rows)
    again.mapping = importer.default_mapping(header, rows, again.sheet)
    again.save()
    assert {p.action for p in importer.build_plan(again)} == {"skip"}
    assert importer.run_import(again, None)["skip"] == 6
    assert Registration.objects.count() == 6


def test_graduation_without_supervisor_gets_follow_up(reference, make_user):
    rows = [("Hanna Smit", 666666, "Tourism", "Erik", "YES", None, "Hotel app", datetime(2027, 2, 1), "?", None, None, GREEN)]
    header, parsed = importer.parse_sheet(io.BytesIO(workbook_bytes(rows)), "2026-2027")
    b = ImportBatch.objects.create(filename="x.xlsx", sheet="2026-2027", header=header, rows=parsed)
    b.mapping = importer.default_mapping(header, parsed, b.sheet)
    b.save()
    importer.run_import(b, None)
    assert GraduationTrack.objects.get().approval == "not_yet"
    assert FollowUp.objects.get().auto_reason == "graduation_approval_missing"


# -- the screens ------------------------------------------------------------------------------------

@pytest.fixture
def admin_client(client, make_user):
    client.force_login(make_user("boss", "Admin"))
    return client


def test_import_is_admin_only(client, make_user, reference):
    client.force_login(make_user("coachy", "Coach"))
    assert client.get("/staff/import/").status_code == 403


def test_full_flow_through_the_screens(admin_client):
    upload = io.BytesIO(workbook_bytes())
    upload.name = "Coach_overview.xlsx"
    response = admin_client.post("/staff/import/", {"file": upload})
    batch = ImportBatch.objects.get()
    assert response.url == f"/staff/import/{batch.pk}/sheet/" and batch.file_data

    response = admin_client.post(f"/staff/import/{batch.pk}/sheet/", {"sheet": "2026-2027"})
    assert response.url == f"/staff/import/{batch.pk}/map/"
    batch.refresh_from_db()
    assert batch.file_data is None and len(batch.rows) == 6  # only the chosen sheet is kept

    page = admin_client.get(f"/staff/import/{batch.pk}/map/").content.decode()
    assert "Buas Emplyee" in page and "Niki" in page

    # Map Niki to Joyce, then preview and import.
    niki_index = list(batch.mapping["coaches"]).index("Niki")
    data = {f"col_{t}": ("" if i is None else i) for t, i in batch.mapping["columns"].items()}
    data["default_date"] = "2026-09-01"
    for i, (value, choice) in enumerate(batch.mapping["programs"].items()):
        data[f"program_domain_{i}"], data[f"program_year_{i}"] = choice["domain"], choice["year"]
    for i, (value, choice) in enumerate(batch.mapping["coaches"].items()):
        data[f"coach_{i}"] = str(Coach.objects.get(first_name="Joyce").pk) if i == niki_index else choice
    for i, (value, choice) in enumerate(batch.mapping["goc"].items()):
        data[f"goc_{i}"] = choice
    for i, (key, c) in enumerate(batch.mapping["colours"].items()):
        data[f"colour_stage_{i}"], data[f"colour_tag_{i}"] = c["stage"], c["tag"]
    data["next"] = "preview"
    response = admin_client.post(f"/staff/import/{batch.pk}/map/", data)
    assert response.url == f"/staff/import/{batch.pk}/preview/"

    preview = admin_client.get(f"/staff/import/{batch.pk}/preview/")
    assert preview.context["counts"]["new"] == 6 and "Import 6 rows" in preview.content.decode()

    admin_client.post(f"/staff/import/{batch.pk}/run/")
    assert Startup.objects.get(description="Bike tours").assigned_coach.first_name == "Joyce"


def test_upload_rejects_non_excel(admin_client):
    bad = io.BytesIO(b"not excel")
    bad.name = "data.xlsx"
    response = admin_client.post("/staff/import/", {"file": bad}, follow=True)
    assert "could not be read" in response.content.decode()
    assert not ImportBatch.objects.exists()


def test_discard(admin_client):
    upload = io.BytesIO(workbook_bytes())
    upload.name = "x.xlsx"
    admin_client.post("/staff/import/", {"file": upload})
    batch = ImportBatch.objects.get()
    admin_client.post(f"/staff/import/{batch.pk}/discard/")
    assert not ImportBatch.objects.exists()
