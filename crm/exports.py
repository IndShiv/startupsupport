"""CSV and Excel exports of the (filtered) student and startup lists.

Column definitions live here so both formats always contain the same data. Text that
came from the public form is neutralised against spreadsheet formula injection.
"""

import csv
import io
from datetime import date, datetime

from django.db.models import Max, Min, Prefetch
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Founder, GraduationTrack

# Spreadsheet apps execute CSV cells starting with these as formulas (CSV injection).
CSV_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_text(value, xlsx=False):
    """Neutralise text a spreadsheet would execute as a formula.

    In .xlsx files only a leading "=" makes a cell a formula (other cells are stored as
    text), so phone numbers like "+31 6…" are left alone there.
    """
    if not isinstance(value, str):
        return value
    prefixes = ("=",) if xlsx else CSV_FORMULA_PREFIXES
    return "'" + value if value.startswith(prefixes) else value


def yes_no(value):
    return {True: "Yes", False: "No"}.get(value, "")


def join(values):
    return "; ".join(v for v in values if v)


def local_dt(value):
    return timezone.localtime(value).replace(tzinfo=None) if value else None


# -- column definitions: (header, width, function(obj) -> value) --------------------------

STUDENT_COLUMNS = [
    ("Student number", 14, lambda s: s.student_number),
    ("First name", 16, lambda s: s.first_name),
    ("Last name", 20, lambda s: s.last_name),
    ("Email", 30, lambda s: s.email),
    ("Phone", 18, lambda s: s.phone),
    ("Domain", 20, lambda s: s.domain.name),
    ("Study year", 24, lambda s: s.study_year.name),
    ("Graduation track", 10, lambda s: yes_no(s.study_year.is_graduation_track)),
    ("Ideas", 8, lambda s: len(s.export_startups)),
    ("Startups / ideas", 40, lambda s: join(st.display_name for st in s.export_startups)),
    ("Stages", 24, lambda s: join(sorted({st.stage.name for st in s.export_startups}))),
    ("Coaches", 24, lambda s: join(sorted({st.assigned_coach.full_name for st in s.export_startups if st.assigned_coach}))),
    ("First registration", 14, lambda s: local_dt(s.first_registration)),
    ("Last registration", 14, lambda s: local_dt(s.last_registration)),
    ("Archived", 10, lambda s: yes_no(bool(s.archived_at))),
]

STARTUP_COLUMNS = [
    ("ID", 7, lambda st: st.pk),
    ("Startup name", 28, lambda st: st.name),
    ("Description", 50, lambda st: st.description),
    ("Goals", 40, lambda st: st.goals),
    ("Stage", 18, lambda st: st.stage.name),
    ("Coach", 20, lambda st: st.assigned_coach.full_name if st.assigned_coach else ""),
    ("Founders", 28, lambda st: join(f.student.full_name for f in st.export_founders)),
    ("Founder emails", 32, lambda st: join(f.student.email for f in st.export_founders)),
    ("Founder student numbers", 16, lambda st: join(f.student.student_number for f in st.export_founders)),
    ("Domains", 20, lambda st: join(sorted({f.student.domain.name for f in st.export_founders}))),
    ("Study years", 24, lambda st: join(sorted({f.student.study_year.name for f in st.export_founders}))),
    ("Paying customers", 10, lambda st: yes_no(st.has_paying_customers)),
    ("Idea validated", 10, lambda st: yes_no(st.idea_validated)),
    ("Tags", 24, lambda st: join(t.name for t in st.tags.all())),
    ("KvK number", 12, lambda st: st.kvk_number),
    ("Graduation track", 10, lambda st: yes_no(bool(st.export_tracks))),
    ("Programme approval", 14, lambda st: join(t.get_approval_display() for t in st.export_tracks)),
    ("Hand-in date", 12, lambda st: st.export_tracks[0].hand_in_date if st.export_tracks else None),
    ("First registration", 14, lambda st: local_dt(st.first_registration)),
    ("Intake held on", 12, lambda st: st.intake_held),
    ("Last activity", 14, lambda st: local_dt(st.last_activity_at)),
    ("Archived", 10, lambda st: yes_no(bool(st.archived_at))),
]


def prepare_students(qs):
    from .models import Startup

    return qs.select_related("domain", "study_year").annotate(
        first_registration=Min("registrations__submitted_at"),
        last_registration=Max("registrations__submitted_at"),
    ).prefetch_related(
        Prefetch("startups", queryset=Startup.objects.select_related("stage", "assigned_coach"), to_attr="export_startups")
    ).order_by("last_name", "first_name")


def prepare_startups(qs):
    return qs.select_related("stage", "assigned_coach").annotate(
        first_registration=Min("registrations__submitted_at"),
        intake_held=Max("registrations__intake_held_on"),
    ).prefetch_related(
        "tags",
        Prefetch("founder_links", queryset=Founder.objects.select_related("student__domain", "student__study_year"), to_attr="export_founders"),
        Prefetch("graduation_tracks", queryset=GraduationTrack.objects.order_by("hand_in_date"), to_attr="export_tracks"),
    ).order_by("-created_at")


def rows(objects, columns):
    for obj in objects:
        yield [fn(obj) for _, _, fn in columns]


# -- writers -----------------------------------------------------------------------------

def to_csv(objects, columns):
    """Comma-separated, UTF-8 with BOM. Returns the file contents as bytes."""
    buffer = io.StringIO()
    buffer.write("﻿")  # BOM, so Excel opens UTF-8 (accents in names) correctly
    writer = csv.writer(buffer)
    writer.writerow([header for header, _, _ in columns])
    for row in rows(objects, columns):
        writer.writerow([_csv_value(v) for v in row])
    return buffer.getvalue().encode("utf-8")


def _csv_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return safe_text(value)


def to_xlsx(objects, columns, *, title, about):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="00426E")
    ws.append([header for header, _, _ in columns])
    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    count = 0
    for row in rows(objects, columns):
        ws.append([safe_text(v, xlsx=True) for v in row])
        count += 1
    for index, (_, width, _) in enumerate(columns, start=1):
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = width
        for cell in ws[letter][1:]:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30

    info = wb.create_sheet("About this export")
    for label, value in about + [("Rows", count)]:
        info.append([label, safe_text(value, xlsx=True)])
    info.append([])
    info.append(["Privacy", "This file contains personal data of BUas students. Store it only in BUas systems, "
                            "share it only with BUSS staff, and delete it when you no longer need it."])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 100
    for label_cell, value_cell in info.iter_rows(min_col=1, max_col=2):
        label_cell.font = Font(bold=True)
        if isinstance(value_cell.value, datetime):
            value_cell.number_format = "yyyy-mm-dd hh:mm"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), count
