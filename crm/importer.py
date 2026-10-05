"""Import historical registrations from a spreadsheet (BUSS coach overview or a Microsoft Forms export).

Flow: parse_workbook() -> guess_columns() / guess_values() -> build_plan() (preview) -> run_import().
Everything before run_import() is read-only, so the preview shows exactly what will happen.
"""

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from siteconfig.models import Domain, PipelineStage, StudyYear

from .models import Activity, Coach, Founder, FollowUp, GraduationTrack, Registration, Startup, Student, Tag
from .services import find_existing_student

MAX_ROWS = 2000
UNKNOWN_DOMAIN = "Unknown (imported)"
UNKNOWN_YEAR = "Unknown (imported)"
SKIP = "__skip__"

# Target fields and the header words that suggest them (Dutch and English).
TARGETS = {
    "full_name": ("Full name", ["student", "full name", "naam", "name"]),
    "first_name": ("First name", ["first name", "voornaam"]),
    "last_name": ("Last name", ["last name", "achternaam", "surname"]),
    "student_number": ("Student number", ["student number", "studentnummer", "studentnumber"]),
    "email": ("Email", ["email", "e-mail", "mail"]),
    "phone": ("Phone", ["phone", "telefoon", "mobile"]),
    "program": ("Study programme / domain", ["studyprogram", "study program", "studyprogramme", "domain", "opleiding", "academy"]),
    "study_year": ("Study year", ["study year", "studiejaar", "year"]),
    "coach": ("Coach", ["buss coach", "coach"]),
    "goc": ("Graduating within own company", ["graduation own company", "goc", "graduating within", "afstuderen"]),
    "company_name": ("Company name", ["company name", "bedrijfsnaam", "startup name"]),
    "description": ("Business (idea) description", ["description", "describe your business", "business idea", "omschrijving"]),
    "goals": ("Goals / how can we help", ["goals", "how can we help"]),
    "paying": ("Paying customers", ["paying customers"]),
    "validated": ("Idea validated", ["validated"]),
    "hand_in": ("Thesis / hand-in deadline", ["deadline thesis", "hand-in", "hand in", "inleverdatum"]),
    "supervisor": ("Graduation supervisor", ["supervisor", "begeleider"]),
    "notes": ("Coaching status / notes", ["status", "notes", "opmerkingen", "comments"]),
    "submitted_at": ("Registration date", ["completion time", "start time", "submitted", "registration date", "date registered"]),
    "response_id": ("Forms response ID", ["id"]),
}

# Colour legend of the BUSS coach overview (fill of the name cell).
COLOUR_NAMES = {
    "FF92D050": "green (active)", "FF00B050": "green (active)",
    "FFFFC000": "orange (inactive)", "FFFFFF00": "yellow",
    "FFFF0000": "red (finished)", "FFC00000": "red (finished)",
    "FF00B0F0": "blue (employee)", "FF0070C0": "blue (employee)",
    "": "no colour",
}
DEFAULT_COLOUR_STAGE = {
    "green (active)": ("Coaching", ""),
    "orange (inactive)": ("Coaching", "inactive (imported)"),
    "red (finished)": ("Alumni", ""),
    "blue (employee)": ("Coaching", ""),
    "no colour": ("Coaching", "status unknown (imported)"),
}

# Separators between team members in one name cell ("Anna Smit & Bram de Vries"). Not commas:
# those are ambiguous ("Smit, Anna"), so such rows get a warning instead.
TEAM_SEPARATORS = re.compile(r"\s+(?:&|/|\+|en|and)\s+|\s*[&/+]\s*", re.IGNORECASE)
NUMBER_RE = re.compile(r"(?<!\d)\d{6}(?!\d)")


def normalise(text):
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9& ]", " ", text.lower())).strip()


# -- parsing ---------------------------------------------------------------------------------

def _cell_value(value):
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value).strip()


def _fill(cell):
    fill = getattr(cell, "fill", None)
    if not fill or fill.fill_type in (None, "none"):
        return ""
    rgb = fill.fgColor.rgb if fill.fgColor is not None else None
    return rgb if isinstance(rgb, str) else ""


def sheet_names(file):
    wb = load_workbook(file, read_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def parse_sheet(file, sheet):
    """Header plus rows ({"row": n, "cells": [...], "fill": "FF92D050"}), skipping empty rows."""
    wb = load_workbook(file, data_only=True)
    ws = wb[sheet]
    header, rows = None, []
    for row in ws.iter_rows():
        values = [_cell_value(c.value) for c in row]
        if header is None:
            if any(values):
                header = [v or f"Column {i + 1}" for i, v in enumerate(values)]
            continue
        if not any(values):
            continue
        rows.append({"row": row[0].row, "cells": values, "fill": _fill(row[0])})
        if len(rows) > MAX_ROWS:
            raise ValueError(f"The sheet has more than {MAX_ROWS} rows; split it up first.")
    return header or [], rows


# -- guessing ----------------------------------------------------------------------------------

def guess_columns(header):
    """Map each target field to the best-matching column index (or None)."""
    mapping, used = {}, set()
    scored = []
    for target, (_label, words) in TARGETS.items():
        for index, title in enumerate(header):
            norm = normalise(title)
            for word in words:
                if (word == norm) or (len(word) > 2 and word in norm):
                    scored.append((len(word) + (10 if word == norm else 0), target, index))
    for _score, target, index in sorted(scored, reverse=True):
        if target not in mapping and index not in used:
            mapping[target] = index
            used.add(index)
    # A single "name" column holds the full name unless separate first/last name columns exist.
    if "first_name" in mapping and "last_name" in mapping:
        mapping.pop("full_name", None)
    return {target: mapping.get(target) for target in TARGETS}


def distinct(rows, index):
    values = {}
    if index is None:
        return values
    for r in rows:
        value = r["cells"][index] if index < len(r["cells"]) else ""
        values[value] = values.get(value, 0) + 1
    return dict(sorted(values.items(), key=lambda kv: (-kv[1], kv[0])))


def guess_program(value):
    """(domain name, study year name or "") for a free-text study programme."""
    n = normalise(value)
    domains = {normalise(d.name): d.name for d in Domain.objects.all()}
    year = ""
    if not n:
        return UNKNOWN_DOMAIN, ""
    if "employ" in n or "emplyee" in n or "medewerker" in n:
        employee = Domain.objects.filter(is_employee=True).first()
        return (employee.name if employee else UNKNOWN_DOMAIN), "Not applicable"
    if "alumn" in n:
        year = "Not applicable"
    if "master" in n:
        year = "BUas Master"
    compact = n.replace(" ", "")
    for norm, name in domains.items():
        if norm == n or norm.replace(" ", "") == compact:
            return name, year
    for norm, name in sorted(domains.items(), key=lambda kv: -len(kv[0])):
        if norm.split(" ")[0] in n.split(" "):
            return name, year
    if "sem" in n.split(" ") or "event" in n:
        return "Leisure & Events", year
    if "mmi" in n.split(" ") or "media" in n:
        return "Media", year
    return UNKNOWN_DOMAIN, year


def guess_coach(value):
    n = normalise(value)
    if not n:
        return ""
    for coach in Coach.objects.all():
        if n in (normalise(coach.full_name), normalise(coach.first_name)):
            return str(coach.pk)
    return ""


def guess_yes_no(value):
    n = normalise(value)
    if n.startswith(("y", "ja", "j ")) or n == "j":
        return "yes"
    if n.startswith(("n", "nee")):
        return "no"
    return "unknown"


def default_mapping(header, rows, sheet):
    columns = guess_columns(header)
    mapping = {"columns": columns, "default_date": default_date_for(sheet).isoformat()}
    mapping["programs"] = {v: dict(zip(("domain", "year"), guess_program(v))) for v in distinct(rows, columns["program"])}
    mapping["coaches"] = {v: guess_coach(v) for v in distinct(rows, columns["coach"])}
    mapping["goc"] = {v: guess_yes_no(v) for v in distinct(rows, columns["goc"])}
    colours = {}
    name_col = columns["full_name"] if columns["full_name"] is not None else columns["first_name"]
    if name_col is not None:
        for r in rows:
            label = COLOUR_NAMES.get(r["fill"], f"colour {r['fill']}")
            stage, tag = DEFAULT_COLOUR_STAGE.get(label, ("Coaching", ""))
            colours.setdefault(r["fill"], {"label": label, "stage": stage, "tag": tag, "count": 0})["count"] += 1
    mapping["colours"] = colours
    return mapping


def default_date_for(sheet):
    """1 September of the academic year in a sheet name like "2026-2027"; otherwise today."""
    match = re.search(r"(20\d\d)\s*[-/]\s*(20\d\d)", sheet or "")
    return date(int(match.group(1)), 9, 1) if match else timezone.localdate()


# -- preview -----------------------------------------------------------------------------------

@dataclass
class PlannedRow:
    row: int
    import_key: str
    people: list = field(default_factory=list)  # [{first_name, last_name, student_number, existing: Student|None}]
    email: str = ""
    phone: str = ""
    domain: str = ""
    year: str = ""
    coach: object = None
    stage: str = ""
    tags: list = field(default_factory=list)
    company_name: str = ""
    description: str = ""
    goals: str = ""
    paying: object = None
    validated: object = None
    goc: str = "unknown"
    hand_in: object = None
    supervisor: str = ""
    notes: str = ""
    submitted_at: object = None
    date_estimated: bool = False
    action: str = "new"  # new | link | skip | error
    warnings: list = field(default_factory=list)
    original: dict = field(default_factory=dict)


def _get(r, columns, target):
    index = columns.get(target)
    if index is None or index >= len(r["cells"]):
        return ""
    return r["cells"][index]


def _parse_date(text):
    if not text:
        return None
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y %H:%M:%S", "%m/%d/%y %H:%M:%S", "%d-%m-%Y %H:%M"):
        try:
            return datetime.strptime(text.strip(), fmt)
        except ValueError:
            continue
    return None


def _split_name(full):
    parts = full.split()
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def _people(r, columns, warnings):
    numbers = NUMBER_RE.findall(_get(r, columns, "student_number"))
    if columns.get("first_name") is not None:
        names = [(_get(r, columns, "first_name"), _get(r, columns, "last_name"))]
    else:
        full = re.sub(r"\(.*?\)", "", _get(r, columns, "full_name")).strip()
        pieces = [p.strip() for p in TEAM_SEPARATORS.split(full) if p.strip()]
        if 1 < len(pieces) <= 4 and all(1 <= len(p.split()) <= 4 for p in pieces):
            names = [_split_name(p) for p in pieces]
            warnings.append(f"Team of {len(pieces)}: one startup with {len(pieces)} founders.")
        else:
            names = [_split_name(full)]
            if len(full.split()) > 4:
                warnings.append("Long name – check that first and last name are split correctly.")
        if "," in full:
            warnings.append("Name contains a comma – check the names after import.")
    if numbers and len(numbers) != len(names):
        warnings.append(f"{len(numbers)} student number(s) for {len(names)} name(s); numbers are given in order.")
    raw_number = _get(r, columns, "student_number")
    if raw_number and not numbers:
        warnings.append(f"No 6-digit student number in “{raw_number}”.")
    people = []
    for i, (first, last) in enumerate(names):
        people.append({"first_name": first.strip()[:100], "last_name": last.strip()[:100], "student_number": numbers[i] if i < len(numbers) else ""})
    return people


def build_plan(batch):
    m = batch.mapping
    columns = {k: (int(v) if v not in (None, "") else None) for k, v in m["columns"].items()}
    coaches = {str(c.pk): c for c in Coach.objects.all()}
    default_dt = timezone.make_aware(datetime.combine(date.fromisoformat(m["default_date"]), time(12)))
    seen_keys = set()
    plan = []
    for r in batch.rows:
        warnings = []
        people = _people(r, columns, warnings)
        first = people[0] if people else {"first_name": "", "last_name": "", "student_number": ""}
        response_id = _get(r, columns, "response_id")
        key_part = response_id or first["student_number"] or normalise(f"{first['first_name']} {first['last_name']}").replace(" ", "-")
        import_key = f"import:{normalise(batch.sheet).replace(' ', '-')}:{key_part}"[:120]
        p = PlannedRow(row=r["row"], import_key=import_key, people=people, warnings=warnings,
                       original={batch.header[i]: v for i, v in enumerate(r["cells"]) if i < len(batch.header) and v})

        p.email = _get(r, columns, "email").lower()
        p.phone = _get(r, columns, "phone")[:20]
        program = m["programs"].get(_get(r, columns, "program"), {"domain": UNKNOWN_DOMAIN, "year": ""})
        p.domain = program["domain"]
        p.goc = m["goc"].get(_get(r, columns, "goc"), "unknown") if columns.get("goc") is not None else "unknown"
        p.year = program.get("year") or _get(r, columns, "study_year") or ""
        if p.goc == "yes":
            p.year = StudyYear.objects.filter(is_graduation_track=True).values_list("name", flat=True).first() or p.year
        if not p.year or not StudyYear.objects.filter(name=p.year).exists():
            p.year = UNKNOWN_YEAR
        coach_value = _get(r, columns, "coach")
        p.coach = coaches.get(m["coaches"].get(coach_value, ""))
        if coach_value and not p.coach:
            p.warnings.append(f"Coach “{coach_value}” is not mapped; no coach assigned.")
        colour = m.get("colours", {}).get(r["fill"]) or {"stage": "Coaching", "tag": ""}
        p.stage = colour["stage"]
        p.tags = [t for t in [colour.get("tag", "")] if t]
        p.company_name = _get(r, columns, "company_name")[:200]
        p.description = _get(r, columns, "description") or "(no description in the imported file)"
        p.goals = _get(r, columns, "goals")
        p.paying = {"yes": True, "no": False}.get(guess_yes_no(_get(r, columns, "paying"))) if columns.get("paying") is not None else None
        p.validated = {"yes": True, "no": False}.get(guess_yes_no(_get(r, columns, "validated"))) if columns.get("validated") is not None else None
        hand_in = _parse_date(_get(r, columns, "hand_in"))
        p.hand_in = hand_in.date() if hand_in else None
        p.supervisor = _get(r, columns, "supervisor").strip(" ?")
        if p.goc == "yes" and not p.hand_in:
            p.warnings.append("Graduating within own company, but no valid hand-in date: tagged for follow-up instead of creating graduation details.")
            p.tags.append("graduation details missing")
        p.notes = _get(r, columns, "notes")
        submitted = _parse_date(_get(r, columns, "submitted_at"))
        p.submitted_at = timezone.make_aware(submitted) if submitted and timezone.is_naive(submitted) else (submitted or default_dt)
        p.date_estimated = submitted is None
        if p.domain == UNKNOWN_DOMAIN:
            p.warnings.append("Study programme unknown: domain set to “Unknown (imported)”.")

        if not first["first_name"]:
            p.action = "error"
            p.warnings.append("No name: row skipped.")
        elif import_key in seen_keys or Registration.objects.filter(forms_response_id=import_key).exists():
            p.action = "skip"
            p.warnings.insert(0, "Already imported (same row key): skipped.")
        else:
            for person in people:
                person["existing"] = find_existing_student(person["student_number"], p.email if len(people) == 1 else "")
            if any(person["existing"] for person in people):
                p.action = "link"
        seen_keys.add(import_key)
        plan.append(p)
    return plan


# -- import ------------------------------------------------------------------------------------

def _unknown(model, name):
    obj, _ = model.objects.get_or_create(name=name, defaults={"active": False, "order": 999})
    return obj


@transaction.atomic
def run_import(batch, user):
    plan = build_plan(batch)
    stages = {s.name: s for s in PipelineStage.objects.all()}
    counts = {"new": 0, "link": 0, "skip": 0, "error": 0, "students_created": 0}
    for p in plan:
        counts[p.action] += 1
        if p.action in ("skip", "error"):
            continue
        domain = Domain.objects.filter(name=p.domain).first() or _unknown(Domain, UNKNOWN_DOMAIN)
        year = StudyYear.objects.filter(name=p.year).first() or _unknown(StudyYear, UNKNOWN_YEAR)

        students = []
        for person in p.people:
            student = person.get("existing")
            if student is None:
                student = Student.objects.create(
                    first_name=person["first_name"], last_name=person["last_name"], student_number=person["student_number"],
                    email=p.email if len(p.people) == 1 else "", phone=p.phone if len(p.people) == 1 else "",
                    domain=domain, study_year=year,
                )
                Student.objects.filter(pk=student.pk).update(created_at=p.submitted_at)
                counts["students_created"] += 1
            students.append(student)

        startup = Startup.objects.create(
            name=p.company_name, description=p.description, goals=p.goals,
            stage=stages.get(p.stage) or PipelineStage.initial(), assigned_coach=p.coach,
            has_paying_customers=p.paying, idea_validated=p.validated,
        )
        for name in p.tags:
            startup.tags.add(Tag.objects.get_or_create(name=name[:50])[0])
        for student in students:
            Founder.objects.get_or_create(startup=startup, student=student, defaults={"joined_on": timezone.localdate(p.submitted_at)})

        registration = Registration.objects.create(
            student=students[0], startup=startup, submitted_at=p.submitted_at, submitted_at_estimated=p.date_estimated,
            source=Registration.Source.IMPORT, forms_response_id=p.import_key, created_by=user,
            is_duplicate_student=p.action == "link",
            answers={"import": {"file": batch.filename, "sheet": batch.sheet, "row": p.row, "values": p.original}},
        )
        if p.goc == "yes" and p.hand_in:
            approved = bool(p.supervisor)
            track = GraduationTrack.objects.create(
                registration=registration, student=students[0], startup=startup,
                approval=GraduationTrack.Approval.YES if approved else GraduationTrack.Approval.NOT_YET,
                topic="(topic not in the imported file)", supervisor_name=p.supervisor or "(unknown)", hand_in_date=p.hand_in,
            )
            if not approved:
                FollowUp.objects.create(
                    student=students[0], startup=startup, due_date=timezone.localdate(),
                    title=f"Check programme approval and supervisor for {students[0].full_name} (imported)",
                    auto_reason="graduation_approval_missing", created_by=user,
                )
        if p.notes:
            Activity.objects.create(startup=startup, kind=Activity.Kind.NOTE, date=timezone.localdate(p.submitted_at),
                                    body=f"Status from {batch.filename} ({batch.sheet}): {p.notes}")
        # Imported notes are not new contact with the startup: keep "last activity" at the registration date.
        Startup.objects.filter(pk=startup.pk).update(created_at=p.submitted_at, last_activity_at=p.submitted_at if p.notes else None)

    batch.status = batch.Status.IMPORTED
    batch.imported_at = timezone.now()
    batch.result = counts
    batch.rows = []  # personal data is not kept in the batch after import
    batch.save()
    return counts
